import openpyxl, collections, csv, json, re, unicodedata
# --- font reading table (bpmfvs phonic_table_Z.txt): char -> [readings], index0 = default, index k -> IVS U+E01E0+k
font = {}
for line in open('phonic_table_Z.txt', encoding='utf-8'):
    p = line.rstrip('\n').split('\t')
    if len(p) < 4: continue
    font[p[0]] = [x.strip() for x in p[3:] if x.strip()]
# only chars actually in shipped font cmap
from fontTools.ttLib import TTFont
tt = TTFont('/Users/jc/Claude/github/codebase/frontend/monorepo/apps/student/public/fonts/BpmfZihiKaiStd-Regular.ttf')
cmap = tt['cmap'].getBestCmap()
uvs = {}
for t in tt['cmap'].tables:
    if t.format == 14:
        for sel, lst in t.uvsDict.items():
            for (u, g) in lst:
                uvs.setdefault(u, set()).add(sel)
def norm(z):
    z = unicodedata.normalize('NFC', z).replace('　',' ').strip()
    return [s for s in z.split(' ') if s]
wb = openpyxl.load_workbook('concised/dict_concised_2014_20260626.xlsx', read_only=True)
allrows = list(wb.worksheets[0].iter_rows(values_only=True)); hdr=allrows[0]; rows=allrows[1:]
VAR_COL = next((i for i,h in enumerate(hdr) if h and '變體注音' in str(h)), None); print('VAR_COL', VAR_COL, hdr[VAR_COL] if VAR_COL is not None else None)
han = re.compile(r'^[㐀-䶿一-鿿\U00020000-\U0002ffff]+$')
# group multi-char words by 字詞名 to detect multi-reading words (多音排序>0 on multi-char)
byword = collections.defaultdict(list)
for r in rows:
    w = (r[0] or '').strip(); z = r[6] or ''
    byword[w].append((r[5], norm(z), (r[7] or "").strip(), norm(str(r[VAR_COL] or "")) if VAR_COL is not None else []))
stats = collections.Counter(); mism = []; charhits = collections.Counter(); charhits_words = collections.defaultdict(list)
fixable_words = 0; unfixable_words = 0; ambiguous = 0
skipped_notinfont = 0; skipped_len = 0
for w, variants in byword.items():
    if not han.match(w) or len(w) < 2:
        continue
    readings = sorted(variants, key=lambda v: (v[0] if isinstance(v[0], int) else 99))
    # dictionary readings per word: list of syllable lists
    dict_reads = [v[1] for v in readings if len(v[1]) == len(w)]
    if not dict_reads:
        skipped_len += 1; continue
    if any(ch not in font or ord(ch) not in cmap for ch in w):
        skipped_notinfont += 1; continue
    stats['words_checked'] += 1
    font_default = [font[ch][0] for ch in w]
    # if font default matches ANY dictionary reading of this word -> ok
    if any(font_default == dr for dr in dict_reads):
        stats['ok'] += 1; continue
    var_reads = [v[3] for v in readings if len(v[3]) == len(w)]
    if any(font_default == vr for vr in var_reads):
        stats['ok_only_via_變體注音'] += 1
    if len(dict_reads) > 1:
        ambiguous += 1  # word has several readings and font default matches none
    dr = dict_reads[0]  # primary (多音排序 lowest first in file order)
    diffs = []
    word_fixable = True
    for i, ch in enumerate(w):
        if font_default[i] != dr[i]:
            if dr[i] in font[ch]:
                k = font[ch].index(dr[i]); diffs.append((i, ch, font_default[i], dr[i], f'IVS+{k}'))
            else:
                diffs.append((i, ch, font_default[i], dr[i], '字型無此讀音')); word_fixable = False
            charhits[ch] += 1
            if len(charhits_words[ch]) < 6: charhits_words[ch].append(w)
    if word_fixable: fixable_words += 1
    else: unfixable_words += 1
    mism.append((w, ' '.join(dr), ' '.join(font_default), '; '.join(f'{c}:{a}→{b}({t})' for _, c, a, b, t in diffs), 'IVS可修' if word_fixable else '字型缺讀音', '多義詞' if len(dict_reads) > 1 else ''))
stats['mismatch_words'] = len(mism); stats['fixable'] = fixable_words; stats['unfixable'] = unfixable_words; stats['ambiguous_multi_reading'] = ambiguous
stats['skipped_len_mismatch'] = skipped_len; stats['skipped_not_in_font'] = skipped_notinfont
stats['distinct_chars'] = len(charhits)
# single-char check: primary reading (多音排序 1 or 0) vs font default
single_bad = []
for w, variants in byword.items():
    if len(w) != 1 or not han.match(w) or w not in font or ord(w) not in cmap: continue
    prim = sorted(variants, key=lambda v: v[0])[0][1]
    if prim and prim[0] != font[w][0]:
        single_bad.append((w, prim[0], font[w][0], 'IVS' if prim[0] in font[w] else '缺'))
stats['single_char_primary_mismatch'] = len(single_bad)
print(json.dumps(stats, ensure_ascii=False, indent=1))
print("top chars:", [(c, n, charhits_words[c][:3]) for c, n in charhits.most_common(40)])
print("single-char primary mismatches sample:", single_bad[:40])
with open('concised_vs_font_mismatch.csv', 'w', newline='', encoding='utf-8-sig') as f:
    wr = csv.writer(f); wr.writerow(['詞目', '簡編本注音', '字型目前顯示', '差異(字:字型→簡編本(修法))', '修法', '備註'])
    for m in mism: wr.writerow(m)
with open('single_char_mismatch.csv', 'w', newline='', encoding='utf-8-sig') as f:
    wr = csv.writer(f); wr.writerow(['字', '簡編本首讀音', '字型預設', '修法']); wr.writerows(single_bad)
# common chars check
for ch in '的了地著得子麼們好不一':
    print(ch, 'font:', font.get(ch), 'dict:', [(v[0], v[1]) for v in byword.get(ch, [])])
