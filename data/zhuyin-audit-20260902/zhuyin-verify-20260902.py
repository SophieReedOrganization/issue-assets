"""End-to-end verification: 簡編本 注音一式 vs what the shipped font renders, per word, per IVS selector.

Ground truth for rendering = the font binary itself (cmap + cmap format 14 + glyf composites),
not the upstream phonic table. Outputs a fix table whose every row is re-decoded and checked.
"""
import collections, csv, json, re, sys, unicodedata
import openpyxl
from fontTools.ttLib import TTFont

FONT = '/Users/jc/Claude/github/codebase/frontend/monorepo/apps/student/public/fonts/BpmfZihiKaiStd-Regular.ttf'
XLSX = 'concised/dict_concised_2014_20260626.xlsx'
IVS_BASE = 0xE01E0

# ---------- pinyin(glyph name) -> zhuyin ----------
INITIALS = [('zh','ㄓ'),('ch','ㄔ'),('sh','ㄕ'),('b','ㄅ'),('p','ㄆ'),('m','ㄇ'),('f','ㄈ'),('d','ㄉ'),('t','ㄊ'),('n','ㄋ'),('l','ㄌ'),
            ('g','ㄍ'),('k','ㄎ'),('h','ㄏ'),('j','ㄐ'),('q','ㄑ'),('x','ㄒ'),('r','ㄖ'),('z','ㄗ'),('c','ㄘ'),('s','ㄙ')]
FINALS = {'a':'ㄚ','o':'ㄛ','e':'ㄜ','eh':'ㄝ','ai':'ㄞ','ei':'ㄟ','ao':'ㄠ','ou':'ㄡ','an':'ㄢ','en':'ㄣ','ang':'ㄤ','eng':'ㄥ','er':'ㄦ',
          'i':'ㄧ','ia':'ㄧㄚ','io':'ㄧㄛ','ie':'ㄧㄝ','iai':'ㄧㄞ','iao':'ㄧㄠ','iu':'ㄧㄡ','iou':'ㄧㄡ','ian':'ㄧㄢ','in':'ㄧㄣ','iang':'ㄧㄤ','ing':'ㄧㄥ','iong':'ㄩㄥ',
          'u':'ㄨ','ua':'ㄨㄚ','uo':'ㄨㄛ','uai':'ㄨㄞ','ui':'ㄨㄟ','uei':'ㄨㄟ','uan':'ㄨㄢ','un':'ㄨㄣ','uen':'ㄨㄣ','uang':'ㄨㄤ','ueng':'ㄨㄥ','ong':'ㄨㄥ',
          'v':'ㄩ','ve':'ㄩㄝ','van':'ㄩㄢ','vn':'ㄩㄣ','vong':'ㄩㄥ',
          'm':'ㄇ','n':'ㄋ','ng':'ㄫ','hm':'ㄏㄇ','hng':'ㄏㄫ'}
# standalone syllables written with y/w
YW = {'yi':'ㄧ','ya':'ㄧㄚ','yo':'ㄧㄛ','ye':'ㄧㄝ','yai':'ㄧㄞ','yao':'ㄧㄠ','you':'ㄧㄡ','yan':'ㄧㄢ','yin':'ㄧㄣ','yang':'ㄧㄤ','ying':'ㄧㄥ','yong':'ㄩㄥ',
      'wu':'ㄨ','wa':'ㄨㄚ','wo':'ㄨㄛ','wai':'ㄨㄞ','wei':'ㄨㄟ','wan':'ㄨㄢ','wen':'ㄨㄣ','wang':'ㄨㄤ','weng':'ㄨㄥ',
      'yu':'ㄩ','yue':'ㄩㄝ','yuan':'ㄩㄢ','yun':'ㄩㄣ'}
TONES = {'1':'','2':'ˊ','3':'ˇ','4':'ˋ'}

def py2zy(name):
    """'z_xing4' -> 'ㄒㄧㄥˋ'; 'z_zi5' -> '˙ㄗ'. Returns None if unparseable."""
    m = re.fullmatch(r'z_([a-z]+)([1-5])', name)
    if not m: return None
    py, tone = m.group(1), m.group(2)
    if py in YW: body = YW[py]
    else:
        ini = ''; rest = py
        for k, v in INITIALS:
            if py.startswith(k): ini, rest = v, py[len(k):]; break
        if ini in ('ㄓ','ㄔ','ㄕ','ㄖ','ㄗ','ㄘ','ㄙ') and rest == 'i': body = ini
        elif ini and rest == '': body = ini            # e.g. z_m? handled by FINALS anyway
        else:
            # j/q/x + u -> ü (must happen before the membership check: 'ue' only exists as 've')
            if ini in ('ㄐ','ㄑ','ㄒ') and rest.startswith('u'):
                alt = 'v' + rest[1:]
                if alt in FINALS: rest = alt
            if rest not in FINALS: return None
            body = ini + FINALS[rest]
    return ('˙' + body) if tone == '5' else body + TONES[tone]

# ---------- decode the font ----------
tt = TTFont(FONT); glyf = tt['glyf']; cmap = tt['cmap'].getBestCmap()
uvs = {}
for t in tt['cmap'].tables:
    if t.format == 14:
        for sel, lst in t.uvsDict.items():
            for (u, g) in lst: uvs.setdefault(u, {})[sel] = g
def zname(gn):
    g = glyf[gn]
    if not g.isComposite(): return None
    zs = [c.glyphName for c in g.components if c.glyphName.startswith('z_')]
    return zs[0] if len(zs) == 1 else None
def is_han(cp): return 0x3400 <= cp <= 0x9FFF or 0x20000 <= cp <= 0x2FFFF
font_read = {}      # char -> list of readings; index k -> selector E01E0+k (k>=1); index 0 default
bad_names = collections.Counter(); bare_not_empty = []
for cp, gn in cmap.items():
    if not is_han(cp): continue
    d = zname(gn)
    if d is None: continue
    dz = py2zy(d)
    if dz is None: bad_names[d] += 1; continue
    reads = [dz]
    vs = uvs.get(cp, {})
    if IVS_BASE in vs and zname(vs[IVS_BASE]) is not None: bare_not_empty.append(chr(cp))
    k = 1
    while IVS_BASE + k in vs:
        zn = zname(vs[IVS_BASE + k]); zz = py2zy(zn) if zn else None
        if zz is None: bad_names[zn or '?'] += 1; break
        reads.append(zz); k += 1
    font_read[chr(cp)] = reads
print('font decoded chars:', len(font_read), '| unparsable glyph names:', dict(bad_names), '| E01E0 not bare:', bare_not_empty[:5])

# cross-check against upstream phonic table (proxy used by audit.py)
table = {}
for line in open('phonic_table_Z.txt', encoding='utf-8'):
    p = line.rstrip('\n').split('\t')
    if len(p) >= 4: table[p[0]] = [x.strip() for x in p[3:] if x.strip()]
tbl_diff = [(c, table[c], font_read[c]) for c in font_read if c in table and table[c] != font_read[c]]
print('table vs font: chars compared', sum(1 for c in font_read if c in table), '| differ:', len(tbl_diff), tbl_diff[:8])

# ---------- dictionary ----------
def norm(z):
    z = unicodedata.normalize('NFC', z or '').replace('　', ' ').strip()
    return [s for s in z.split(' ') if s]
wb = openpyxl.load_workbook(XLSX, read_only=True)
rows = list(wb.worksheets[0].iter_rows(values_only=True))[1:]
HAN = re.compile(r'^[㐀-䶿一-鿿\U00020000-\U0002ffff]+$')
byword = collections.defaultdict(list)
for r in rows:
    w = (r[0] or '').strip()
    byword[w].append({'order': r[5] if isinstance(r[5], int) else 99, 'zy': norm(r[6]), 'vtype': (r[7] or '').strip(), 'vzy': norm(r[8]) if r[8] else [], 'id': r[1], 'gloss': re.sub(r'\s+', ' ', (r[13] or ''))[:60]})

def encode(w, syl):
    """Return (ivs_string, per-char detail, ok) rendering `syl` for word `w` using the font."""
    out = []; detail = []; ok = True
    for ch, s in zip(w, syl):
        reads = font_read.get(ch)
        if reads is None: detail.append((ch, s, 'font-has-no-zhuyin')); ok = False; out.append(ch); continue
        if reads[0] == s: out.append(ch); detail.append((ch, s, 'default'))
        elif s in reads:
            k = reads.index(s); out.append(ch + chr(IVS_BASE + k)); detail.append((ch, s, f'IVS+{k}'))
        else: detail.append((ch, s, f'missing(font:{"/".join(reads)})')); ok = False; out.append(ch)
    return ''.join(out), detail, ok

def decode(s):
    """Simulate what the font shows for string s: list of readings (None for non-han)."""
    res = []; i = 0
    while i < len(s):
        ch = s[i]; i += 1
        sel = None
        if i < len(s) and 0xE0100 <= ord(s[i]) <= 0xE01EF: sel = ord(s[i]); i += 1
        if not is_han(ord(ch)): res.append(None); continue
        reads = font_read.get(ch)
        if reads is None: res.append(None); continue
        if sel is None: res.append(reads[0])
        else:
            k = sel - IVS_BASE
            res.append(reads[k] if 1 <= k < len(reads) else ('' if k == 0 else reads[0]))
    return res

stats = collections.Counter(); fix_rows = []; unfixable = []; skipped = []; multi = []; already_ok = 0
for w, ents in byword.items():
    if len(w) < 2 or not HAN.match(w): continue
    ents = sorted(ents, key=lambda e: (e['vtype'] != '', e['order']))   # main row first, then by 多音排序
    def segment(w, syl):
        # syllable count != char count: try to split the concatenated zhuyin using each char's font readings
        s = ''.join(syl); sols = []
        def rec(i, pos, acc):
            if i == len(w):
                if pos == len(s): sols.append(list(acc))
                return
            for rd in font_read.get(w[i], []):
                if s.startswith(rd, pos): rec(i + 1, pos + len(rd), acc + [rd])
        rec(0, 0, [])
        return sols[0] if len(sols) == 1 else None
    for e in ents:
        if len(e['zy']) != len(w):
            seg = segment(w, e['zy'])
            if seg: e['zy'] = seg; e['segmented'] = True
    valid = [e for e in ents if len(e['zy']) == len(w)]
    if not valid:
        skipped.append((w, ' | '.join(' '.join(e['zy']) for e in ents))); continue
    stats['words'] += 1
    fdef = [font_read.get(ch, [None])[0] for ch in w]
    readings = [e['zy'] for e in valid]
    uniq = []
    for rd in readings:
        if rd not in uniq: uniq.append(rd)
    target = valid[0]['zy']
    if len(uniq) > 1:
        glosses = ' ‖ '.join(f"{' '.join(e['zy'])}＝{e['gloss']}" for e in valid if e['zy'] in uniq and [x['zy'] for x in valid].index(e['zy']) == valid.index(e))
        multi.append((w, ' / '.join(' '.join(u) for u in uniq), ' '.join(x or '?' for x in fdef), '目前顯示=第%d讀音' % (uniq.index(fdef) + 1) if fdef in uniq else '目前顯示都不是', glosses))
    if fdef in uniq:
        already_ok += 1; continue
    if len(uniq) > 1:
        stats['manual_multi'] += 1; continue   # listed in multi-reading csv for human decision
    ivs, detail, ok = encode(w, target)
    if not ok:
        unfixable.append((w, ' '.join(target), ' '.join(x or '?' for x in fdef), '; '.join(f'{c}:{n}' for c, s, n in detail if not n.startswith(('default', 'IVS')))))
        continue
    # verify by decoding
    dec = decode(ivs)
    assert dec == target, (w, dec, target)
    stats['verified_fix'] += 1
    fix_rows.append((w, ' '.join(target), ' '.join(x or '?' for x in fdef), ivs, ' '.join(f'U+{ord(c):04X}' for c in ivs), '; '.join(f'{c}:{n}' for c, s, n in detail if n != 'default'), '多義詞→取簡編本主讀音' if len(uniq) > 1 else ''))
stats['already_ok'] = already_ok; stats['unfixable'] = len(unfixable); stats['skipped_len'] = len(skipped); stats['multi_reading_words'] = len(multi)
print(json.dumps(stats, ensure_ascii=False))
chars = collections.Counter()
for r in fix_rows:
    for part in r[5].split('; '):
        if part: chars[part[0]] += 1
print('distinct chars needing selector:', len(chars), chars.most_common(12))

with open('zhuyin-fix-table-20260902.csv', 'w', newline='', encoding='utf-8-sig') as f:
    wr = csv.writer(f); wr.writerow(['詞目', '簡編本注音一式', '字型目前顯示', '插碼後字串', '碼點', '哪個字用第幾讀音', '備註']); wr.writerows(fix_rows)
with open('zhuyin-fix-table-20260902.json', 'w', encoding='utf-8') as f:
    json.dump({r[0]: r[3] for r in fix_rows}, f, ensure_ascii=False, indent=0)
with open('zhuyin-unfixable-20260902.csv', 'w', newline='', encoding='utf-8-sig') as f:
    wr = csv.writer(f); wr.writerow(['詞目', '簡編本注音一式', '字型目前顯示', '缺什麼']); wr.writerows(unfixable)
with open('zhuyin-multi-reading-words-20260902.csv', 'w', newline='', encoding='utf-8-sig') as f:
    wr = csv.writer(f); wr.writerow(['詞目', '簡編本各讀音（主讀音在前）', '字型目前顯示', '目前狀態', '各讀音釋義']); wr.writerows(multi)
with open('zhuyin-skipped-20260902.csv', 'w', newline='', encoding='utf-8-sig') as f:
    wr = csv.writer(f); wr.writerow(['詞目', '簡編本注音（音節數≠字數）']); wr.writerows(skipped)
print('skipped sample:', skipped[:12])
print('unfixable:', len(unfixable), unfixable[:6])
